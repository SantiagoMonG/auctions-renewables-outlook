from __future__ import annotations

import calendar
import csv
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from src.config import PROJECT_ROOT

SHEET_TO_TECH = {
    "Resumen-DiarioHorario-Eólico": "wind",
    "Resumen-DiarioHorario-Solar": "solar",
}

INPUT_ROOT_CANDIDATES = (
    PROJECT_ROOT / "data" / "raw" / "country_case" / "chile" / "curtailment_chile",
    PROJECT_ROOT / "data" / "raw" / "country_case" / "Chile" / "curtailment-chile",
)

HOURLY_TOTAL_COLUMNS = [
    "date",
    "year",
    "month",
    "day",
    "hour",
    "technology",
    "curtailment_mwh",
    "source_file",
]

HOURLY_BY_PLANT_COLUMNS = [
    "date",
    "year",
    "month",
    "day",
    "hour",
    "technology",
    "plant",
    "curtailment_mwh",
    "source_file",
]


@dataclass
class DailyBlock:
    start_row: int
    total_row: int


@dataclass
class ParserValidationIssue:
    source_file: str
    sheet: str
    issue: str


class WorkbookUnavailableError(RuntimeError):
    pass


def _import_openpyxl() -> object:
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise WorkbookUnavailableError(
            "openpyxl is required for this parser. Install dependencies and rerun."
        ) from exc
    return openpyxl


def parse_numeric(value: object) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    normalized = str(value).strip().replace(".", "").replace(",", ".")
    if not normalized:
        return 0.0
    try:
        return float(normalized)
    except ValueError:
        return 0.0


def get_input_root() -> Path:
    for candidate in INPUT_ROOT_CANDIDATES:
        if candidate.exists():
            return candidate
    raise FileNotFoundError(
        "Could not find curtailment input folder. Checked: "
        + ", ".join(str(path) for path in INPUT_ROOT_CANDIDATES)
    )


def list_monthly_files(input_root: Path) -> list[Path]:
    files = sorted(input_root.glob("*/*.xlsx"))
    filtered = [path for path in files if path.stem.count("-") == 1]
    if not filtered:
        raise FileNotFoundError(f"No monthly files found under: {input_root}")
    return filtered


def parse_year_month(path: Path) -> tuple[int, int]:
    year_str, month_str = path.stem.split("-")
    return int(year_str), int(month_str)


def find_daily_blocks(sheet) -> tuple[list[DailyBlock], list[str]]:
    blocks: list[DailyBlock] = []
    issues: list[str] = []
    in_block = False
    start_row = -1

    for row_idx in range(1, sheet.max_row + 1):
        marker = str(sheet.cell(row=row_idx, column=2).value or "").strip()

        if marker == "Central/Hora":
            if in_block:
                issues.append(f"Nested block start detected at row {row_idx}.")
            in_block = True
            start_row = row_idx
            continue

        if marker == "Total" and in_block:
            blocks.append(DailyBlock(start_row=start_row, total_row=row_idx))
            in_block = False
            start_row = -1

    if in_block:
        issues.append("A block starts with Central/Hora but has no Total row.")

    return blocks, issues


def extract_hourly_values(sheet, row_idx: int) -> list[float]:
    # E:AB == columns 5..28 -> 24 values
    values = [parse_numeric(sheet.cell(row=row_idx, column=col_idx).value) for col_idx in range(5, 29)]
    if len(values) != 24:
        raise ValueError(f"Expected 24 hourly values, found {len(values)} at row {row_idx}.")
    return values


def parse_sheet(
    sheet,
    technology: str,
    source_file: str,
    year: int,
    month: int,
    include_plant_output: bool,
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[str]]:
    issues: list[str] = []
    totals_rows: list[dict[str, object]] = []
    by_plant_rows: list[dict[str, object]] = []

    blocks, block_issues = find_daily_blocks(sheet)
    issues.extend(block_issues)

    expected_days = calendar.monthrange(year, month)[1]
    if len(blocks) != expected_days:
        issues.append(
            f"Detected {len(blocks)} blocks, expected {expected_days} for {year}-{month:02d}."
        )

    for block_idx, block in enumerate(blocks, start=1):
        try:
            day_date = date(year, month, block_idx)
        except ValueError:
            issues.append(f"Invalid day from block index {block_idx}; skipping block.")
            continue

        total_values = extract_hourly_values(sheet, block.total_row)
        for hour, value in enumerate(total_values, start=1):
            totals_rows.append(
                {
                    "date": day_date.isoformat(),
                    "year": year,
                    "month": month,
                    "day": day_date.day,
                    "hour": hour,
                    "technology": technology,
                    "curtailment_mwh": value,
                    "source_file": source_file,
                }
            )

        if not include_plant_output:
            continue

        plant_hourly_sum = np.zeros(24, dtype=float)
        for row_idx in range(block.start_row + 1, block.total_row):
            plant_name = str(sheet.cell(row=row_idx, column=2).value or "").strip()
            if not plant_name:
                continue

            hourly_values = extract_hourly_values(sheet, row_idx)
            plant_hourly_sum += np.array(hourly_values)
            for hour, value in enumerate(hourly_values, start=1):
                by_plant_rows.append(
                    {
                        "date": day_date.isoformat(),
                        "year": year,
                        "month": month,
                        "day": day_date.day,
                        "hour": hour,
                        "technology": technology,
                        "plant": plant_name,
                        "curtailment_mwh": value,
                        "source_file": source_file,
                    }
                )

        total_arr = np.array(total_values)
        if not np.allclose(plant_hourly_sum, total_arr, atol=1e-4):
            issues.append(
                f"Plant sum mismatch for block/day {block_idx}: max hourly abs diff "
                f"{np.max(np.abs(plant_hourly_sum - total_arr)):.6f}."
            )

    return totals_rows, by_plant_rows, issues


def write_csv(path: Path, rows: Sequence[dict[str, object]], columns: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(columns))
        writer.writeheader()
        writer.writerows(rows)


def write_validation_report(path: Path, issues: Iterable[ParserValidationIssue]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    issue_list = list(issues)
    with path.open("w", encoding="utf-8") as f:
        if not issue_list:
            f.write("No parser validation issues found.\n")
            return

        f.write("Parser validation report\n")
        f.write("========================\n\n")
        for item in issue_list:
            f.write(f"file={item.source_file} | sheet={item.sheet} | issue={item.issue}\n")


def make_monthly_aggregate(hourly_total_csv: Path, monthly_csv: Path) -> None:
    df = pd.read_csv(hourly_total_csv)
    df["date"] = pd.to_datetime(df["date"])
    out = (
        df.groupby(["year", "month", "technology"], as_index=False)["curtailment_mwh"]
        .sum()
        .sort_values(["year", "month", "technology"])
    )
    monthly_csv.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(monthly_csv, index=False)


def save_heatmap(df: pd.DataFrame, title: str, output_path: Path) -> None:
    work = df.copy()
    work["date"] = pd.to_datetime(work["date"])
    work["date_label"] = work["date"].dt.strftime("%Y-%m-%d")
    heat = work.pivot_table(
        index="hour",
        columns="date_label",
        values="curtailment_mwh",
        aggfunc="sum",
        fill_value=0.0,
    )
    heat = heat.reindex(index=list(range(1, 25)), fill_value=0.0)

    fig, ax = plt.subplots(figsize=(24, 6), dpi=150)
    im = ax.imshow(heat.values, aspect="auto", origin="lower", cmap="magma")
    ax.set_title(title)
    ax.set_xlabel("Date")
    ax.set_ylabel("Hour")
    ax.set_yticks(np.arange(24))
    ax.set_yticklabels([str(h) for h in range(1, 25)])

    tick_step = max(1, heat.shape[1] // 30)
    tick_positions = np.arange(0, heat.shape[1], tick_step)
    ax.set_xticks(tick_positions)
    ax.set_xticklabels([heat.columns[i] for i in tick_positions], rotation=45, ha="right", fontsize=7)

    cbar = fig.colorbar(im, ax=ax)
    cbar.set_label("Curtailment (MWh)")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.tight_layout()
    fig.savefig(output_path)
    plt.close(fig)


def build_heatmaps(hourly_total_csv: Path, figures_dir: Path) -> None:
    df = pd.read_csv(hourly_total_csv)

    total_df = (
        df.groupby(["date", "hour"], as_index=False)["curtailment_mwh"].sum().assign(technology="total")
    )
    solar_df = df[df["technology"] == "solar"].copy()
    wind_df = df[df["technology"] == "wind"].copy()

    save_heatmap(total_df, "Chile Curtailment Heatmap (Solar + Wind)", figures_dir / "heatmap_total.png")
    save_heatmap(solar_df, "Chile Curtailment Heatmap (Solar)", figures_dir / "heatmap_solar.png")
    save_heatmap(wind_df, "Chile Curtailment Heatmap (Wind)", figures_dir / "heatmap_wind.png")


def run_parser(include_plant_output: bool = True) -> None:
    openpyxl = _import_openpyxl()

    input_root = get_input_root()
    files = list_monthly_files(input_root)

    output_dir = PROJECT_ROOT / "outputs"
    hourly_total_path = output_dir / "hourly_total.csv"
    hourly_by_plant_path = output_dir / "hourly_by_plant.csv"
    validation_path = output_dir / "parser_validation_report.txt"
    monthly_agg_path = output_dir / "monthly_from_hourly.csv"
    figures_dir = output_dir / "figures"

    all_totals: list[dict[str, object]] = []
    all_plants: list[dict[str, object]] = []
    validation_issues: list[ParserValidationIssue] = []

    for file_path in files:
        year, month = parse_year_month(file_path)
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

        for sheet_name, technology in SHEET_TO_TECH.items():
            if sheet_name not in workbook.sheetnames:
                validation_issues.append(
                    ParserValidationIssue(
                        source_file=file_path.name,
                        sheet=sheet_name,
                        issue="Sheet missing",
                    )
                )
                continue

            sheet = workbook[sheet_name]
            try:
                totals, plants, issues = parse_sheet(
                    sheet=sheet,
                    technology=technology,
                    source_file=file_path.name,
                    year=year,
                    month=month,
                    include_plant_output=include_plant_output,
                )
                all_totals.extend(totals)
                all_plants.extend(plants)
                for issue in issues:
                    validation_issues.append(
                        ParserValidationIssue(source_file=file_path.name, sheet=sheet_name, issue=issue)
                    )
            except Exception as exc:  # broad catch to keep full-run reportable per file/sheet
                validation_issues.append(
                    ParserValidationIssue(
                        source_file=file_path.name,
                        sheet=sheet_name,
                        issue=f"Failed to parse sheet cleanly: {exc}",
                    )
                )

        workbook.close()

    write_csv(hourly_total_path, all_totals, HOURLY_TOTAL_COLUMNS)
    if include_plant_output:
        write_csv(hourly_by_plant_path, all_plants, HOURLY_BY_PLANT_COLUMNS)

    write_validation_report(validation_path, validation_issues)
    make_monthly_aggregate(hourly_total_path, monthly_agg_path)
    build_heatmaps(hourly_total_path, figures_dir)

    print(f"Saved: {hourly_total_path}")
    if include_plant_output:
        print(f"Saved: {hourly_by_plant_path}")
    print(f"Saved: {validation_path}")
    print(f"Saved: {monthly_agg_path}")
    print(f"Saved figures in: {figures_dir}")


if __name__ == "__main__":
    run_parser(include_plant_output=True)
